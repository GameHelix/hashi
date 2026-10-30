"""Export the last hours of 15-minute candles on Hyperliquid to Excel: every coin with $1M+ traded, one row per candle, a filter on the header, and the runs of four green candles on a second sheet.

    export_15m.py OUT.xlsx [HOURS]
"""
import json, sys, time, urllib.request
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
HOURS = int(sys.argv[2]) if len(sys.argv) > 2 else 24
API = "https://api.hyperliquid.xyz/info"

def info(body):
    req = urllib.request.Request(API, data=json.dumps(body).encode(), headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.load(r)

meta, ctxs = info({"type": "metaAndAssetCtxs"})
coins = [(a["name"], float(c["dayNtlVlm"])) for a, c in zip(meta["universe"], ctxs) if not a.get("isDelisted") and float(c["dayNtlVlm"]) >= 1e6]
now = int(time.time() * 1000)
candles, runs = [], []
for coin, day_vol in coins:
    ks = info({"type": "candleSnapshot", "req": {"coin": coin, "interval": "15m", "startTime": now - HOURS * 3600 * 1000, "endTime": now}})
    time.sleep(0.6)
    ks = [k for k in ks if k["T"] <= now]  # closed candles only
    rows = []
    for k in ks:
        o, h, l, c, v, n = float(k["o"]), float(k["h"]), float(k["l"]), float(k["c"]), float(k["v"]), int(k["n"])
        rows.append([coin, datetime.fromtimestamp(k["t"] / 1000, timezone.utc).replace(tzinfo=None), o, h, l, c, v, v * c, n, c / o - 1 if o else 0, c > o, False])
    for i in range(3, len(rows)):
        rows[i][11] = all(rows[j][10] for j in range(i - 3, i + 1))
    candles += rows
    for i in range(4, len(rows)):
        if rows[i - 1][11]:  # the four candles before this one were green
            entry = rows[i][2]
            after = rows[i:]
            runs.append([coin, rows[i][1], rows[i - 1][5] / rows[i - 4][2] - 1, entry, max(r[3] for r in after) / entry - 1, min(r[4] for r in after) / entry - 1, after[-1][5] / entry - 1, len(after)])

wb = Workbook()
font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
fill = PatternFill("solid", fgColor="DDEBF7")

def sheet(ws, headers, rows, formats, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = bold, fill, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for col, fmt in formats.items():
        for cell in ws[col][1:]:
            cell.number_format = fmt
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = font
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30

ws = wb.active
ws.title = "candles"
sheet(ws, ["coin", "open time (UTC)", "open", "high", "low", "close", "volume (coin)", "volume (USD, ≈ coin × close)", "trades", "change", "green", "4th green in a row"],
      candles, {"B": "yyyy-mm-dd hh:mm", "C": "0.00####", "D": "0.00####", "E": "0.00####", "F": "0.00####", "G": "#,##0.###", "H": "#,##0", "I": "#,##0", "J": "0.00%"},
      [9, 17, 12, 12, 12, 12, 14, 16, 9, 9, 7, 10])
ws2 = wb.create_sheet("four green runs")
sheet(ws2, ["coin", "entry: open of the next candle (UTC)", "rise over the 4 candles", "entry price", "best after (highest high)", "worst after (lowest low)", "at the last candle", "candles after"],
      sorted(runs, key=lambda r: -r[4]), {"B": "yyyy-mm-dd hh:mm", "C": "0.0%", "D": "0.00####", "E": "0.0%", "F": "0.0%", "G": "0.0%"}, [9, 22, 14, 12, 14, 14, 12, 9])
ws3 = wb.create_sheet("about")
for line in [f"Hyperliquid perpetuals, 15-minute candles, the last {HOURS} hours, closed candles only.",
             f"Coins with at least $1,000,000 traded in the prior day: {len(coins)}. Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC from api.hyperliquid.xyz.",
             "volume (USD) is the coin volume times the close, an approximation. change is close / open - 1.",
             "4th green in a row marks a candle that is the fourth green in a row; the runs sheet buys at the open of the candle after it.",
             "best after and worst after are the highest high and lowest low of all later candles in the window against the entry price."]:
    ws3.append([line])
    ws3[ws3.max_row][0].font = font
ws3.column_dimensions["A"].width = 120
wb.save(OUT)
print(f"{OUT}: {len(candles)} candles of {len(coins)} coins, {len(runs)} four-green runs")
